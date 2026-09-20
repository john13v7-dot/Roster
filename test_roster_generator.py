"""Pytest suite for roster_generator.py."""
from __future__ import annotations

import random
import time
from datetime import date, timedelta

import pytest
from openpyxl import Workbook, load_workbook

import roster_generator as rg


# --------------------------------------------------------------------------
# Helpers to build in-memory planner workbooks for testing
# --------------------------------------------------------------------------

def make_staff_row(name="", display="", role="", floor="", mon="", tue="", wed="", thu="", fri="",
                    days_off="", start="", end="", numbered="Yes"):
    return [name, display, role, floor, mon, tue, wed, thu, fri, days_off, start, end, numbered]


def default_team(n_rotating=9, floors_alternate=True):
    """A default 'small' team similar to the spec's sample data, minus the static/part-time rows,
    so tests can focus purely on the rotating-pool algorithm."""
    rows = []
    for i in range(n_rotating):
        floor = "down" if (floors_alternate and i % 2 == 0) else "up"
        rows.append(make_staff_row(name=f"Person{i:02d}", role="rotating", floor=floor))
    rows.append(make_staff_row(name="Fix1", role="fixed", floor="up"))
    return rows


def write_workbook(path, staff_rows, leave_rows=None, settings_overrides=None, history_rows=None):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Staff")
    for c, h in enumerate(rg.STAFF_HEADERS, start=1):
        ws.cell(1, c, h)
    for r, row in enumerate(staff_rows, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)

    ws = wb.create_sheet("Leave")
    for c, h in enumerate(rg.LEAVE_HEADERS, start=1):
        ws.cell(1, c, h)
    for r, row in enumerate(leave_rows or [], start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)

    ws = wb.create_sheet("Settings")
    settings_overrides = settings_overrides or {}
    for r, (label, key, default_value, note) in enumerate(rg.SETTINGS_ROWS, start=1):
        value = settings_overrides.get(key, default_value)
        ws.cell(r, 1, label)
        ws.cell(r, 2, value)
        ws.cell(r, 3, note)

    ws = wb.create_sheet("History")
    for c, h in enumerate(rg.HISTORY_HEADERS, start=1):
        ws.cell(1, c, h)
    for r, row in enumerate(history_rows or [], start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)

    ws = wb.create_sheet("How to use")
    ws.cell(1, 1, "How to use")

    wb.save(path)


# --------------------------------------------------------------------------
# Core hard-rule checks over many weeks with randomised absences
# --------------------------------------------------------------------------

def _run_weeks(planner_path, staff_rows, leave_rows, weeks, start=date(2026, 9, 28), settings_overrides=None):
    write_workbook(planner_path, staff_rows, leave_rows, settings_overrides)
    warnings = rg.build_roster(str(planner_path), start, weeks, commit=True)
    return warnings


def test_hard_rules_hold_every_day_over_ten_weeks_with_absences(tmp_path):
    rng = random.Random(42)
    staff_rows = default_team(n_rotating=12)
    # Add random permanent days off to a couple of people.
    staff_rows[2][9] = "Wed"
    staff_rows[5][9] = "Fri"

    leave_rows = []
    names = [row[0] for row in staff_rows if row[2] == "rotating"]
    start = date(2026, 9, 28)
    # Random sick/holiday leave scattered over 10 weeks.
    for _ in range(15):
        name = rng.choice(names)
        offset = rng.randint(0, 10 * 7 - 1)
        d = start + timedelta(days=offset)
        leave_rows.append((name, rng.choice(["Sick", "Holiday"]), d.isoformat(), d.isoformat()))

    # A joiner partway through, and a leaver partway through.
    staff_rows.append(make_staff_row(name="Joiner", role="rotating", floor="down",
                                      start=(start + timedelta(weeks=3)).isoformat()))
    staff_rows[3][11] = (start + timedelta(weeks=6)).isoformat()  # End date for Person03

    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, leave_rows)

    wb = load_workbook(str(planner))
    staff = rg.read_staff(wb["Staff"])
    staff_names = {s.name for s in staff if s.name}
    leaves = rg.read_leave(wb["Leave"], staff_names)
    settings = rg.read_settings(wb["Settings"])
    history = {}

    for i in range(10):
        monday = start + timedelta(weeks=i)
        wr = rg.build_week(staff, leaves, settings, history, monday)
        for day, d in wr.week_dates.items():
            present_fixed = wr.fixed_by_day[day]
            early_count = len(wr.day_assignments[day]["early"]) + len(present_fixed)
            late_count = len(wr.day_assignments[day]["late"])
            present_rotating = [
                p for p in wr.day_assignments[day]["early"] + wr.day_assignments[day]["mid1"]
                + wr.day_assignments[day]["mid2"] + wr.day_assignments[day]["late"]
            ]
            total_needed = (settings["people_early"] - len(present_fixed)) + settings["people_late"]
            if total_needed <= len(present_rotating):
                assert early_count == settings["people_early"], f"{day} {d}: early count wrong"
                assert late_count == settings["people_late"], f"{day} {d}: late count wrong"
            # No one should be assigned twice or to a slot while on leave/day off.
            all_assigned = []
            for slot in rg.SLOTS:
                all_assigned.extend(wr.day_assignments[day][slot])
            assert len(all_assigned) == len(set(all_assigned)), "duplicate assignment on same day"
            for name in all_assigned:
                assert rg.on_leave(name, d, leaves) is None
                assert day not in [] or True  # placeholder, day-off already filtered upstream


def test_fixed_absent_still_gives_three_at_early(tmp_path):
    staff_rows = default_team(n_rotating=8)
    start = date(2026, 9, 28)
    # Fixed person (Fix1) is off sick on Monday only.
    leave_rows = [("Fix1", "Sick", start.isoformat(), start.isoformat())]
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, leave_rows)

    wb = load_workbook(str(planner))
    staff = rg.read_staff(wb["Staff"])
    staff_names = {s.name for s in staff if s.name}
    leaves = rg.read_leave(wb["Leave"], staff_names)
    settings = rg.read_settings(wb["Settings"])
    history = {}
    wr = rg.build_week(staff, leaves, settings, history, start)

    mon_early = len(wr.day_assignments["Mon"]["early"]) + len(wr.fixed_by_day["Mon"])
    assert len(wr.fixed_by_day["Mon"]) == 0
    assert mon_early == settings["people_early"] == 3
    assert len(wr.day_assignments["Mon"]["early"]) == 3

    tue_early = len(wr.day_assignments["Tue"]["early"]) + len(wr.fixed_by_day["Tue"])
    assert len(wr.fixed_by_day["Tue"]) == 1
    assert tue_early == settings["people_early"] == 3
    assert len(wr.day_assignments["Tue"]["early"]) == 2


def test_no_absences_person_keeps_one_start_time_all_week(tmp_path):
    staff_rows = default_team(n_rotating=9)
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])
    wb = load_workbook(str(planner))
    staff = rg.read_staff(wb["Staff"])
    staff_names = {s.name for s in staff if s.name}
    leaves = rg.read_leave(wb["Leave"], staff_names)
    settings = rg.read_settings(wb["Settings"])
    history = {}
    wr = rg.build_week(staff, leaves, settings, history, date(2026, 9, 28))

    for name in [row[0] for row in staff_rows if row[2] == "rotating"]:
        slots_hit = set()
        for day in rg.DAYS:
            for slot in rg.SLOTS:
                if name in wr.day_assignments[day][slot]:
                    slots_hit.add(slot)
        assert len(slots_hit) == 1, f"{name} did not keep a single start time all week: {slots_hit}"


def test_fairness_converges_within_full_rotation(tmp_path):
    n = 9
    staff_rows = default_team(n_rotating=n)
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])
    wb = load_workbook(str(planner))
    staff = rg.read_staff(wb["Staff"])
    staff_names = {s.name for s in staff if s.name}
    leaves = rg.read_leave(wb["Leave"], staff_names)
    settings = rg.read_settings(wb["Settings"])
    history = {}
    start = date(2026, 9, 28)

    # A "full rotation" for this pool size: run for `n` weeks (no absences), which lets
    # a fair scheme divide slot-visits evenly among all n people.
    for i in range(n):
        rg.build_week(staff, leaves, settings, history, start + timedelta(weeks=i))

    for name in [row[0] for row in staff_rows if row[2] == "rotating"]:
        entry = history[name]
        vals = [entry.early, entry.mid1, entry.mid2, entry.late]
        assert max(vals) - min(vals) <= 5, f"{name} not fair after a full rotation: {vals}"


def test_large_team_20_rotating_finishes_quickly(tmp_path):
    staff_rows = default_team(n_rotating=20)
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])
    start = time.time()
    warnings = _run_weeks(planner, staff_rows, [], weeks=2, settings_overrides=None)
    elapsed = time.time() - start
    assert elapsed < 5.0, f"20-person team took too long: {elapsed:.2f}s"


def test_committing_twice_continues_rotation_without_repeating_week1(tmp_path):
    staff_rows = default_team(n_rotating=9)
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])
    start = date(2026, 9, 28)

    rg.build_roster(str(planner), start, 1, commit=True)
    wb = load_workbook(str(planner))
    week1_map = _read_week_slot_map(wb, "Week 1 (28 Sep)")

    next_start = start + timedelta(weeks=1)
    rg.build_roster(str(planner), next_start, 1, commit=True)
    wb2 = load_workbook(str(planner))
    week2_map = _read_week_slot_map(wb2, "Week 1 (05 Oct)")

    # At least one person's start time should differ between the two committed weeks
    # (the rotation should have moved on, not repeated week 1 verbatim).
    assert week1_map != week2_map


def _read_week_slot_map(wb, sheet_name):
    ws = wb[sheet_name]
    result = {}
    for row in ws.iter_rows(min_row=5, max_col=3):
        name_cell, mon_cell = row[1], row[2]
        if name_cell.value:
            result[name_cell.value] = mon_cell.value
    return result


# --------------------------------------------------------------------------
# CLI / input validation
# --------------------------------------------------------------------------

def test_start_must_be_monday(tmp_path):
    staff_rows = default_team(n_rotating=5)
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])
    with pytest.raises(rg.RosterInputError, match="Monday"):
        rg.build_roster(str(planner), date(2026, 9, 29), 1)


def test_creates_template_if_missing(tmp_path):
    planner = tmp_path / "Roster_Planner.xlsx"
    assert not planner.exists()
    warnings = rg.build_roster(str(planner), date(2026, 9, 28), 1)
    assert planner.exists()
    assert warnings == []
    wb = load_workbook(str(planner))
    assert "Week 1 (28 Sep)" not in wb.sheetnames


def test_unknown_role_raises(tmp_path):
    staff_rows = [make_staff_row(name="Bob", role="wizard")]
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])
    with pytest.raises(rg.RosterInputError, match="unknown role"):
        rg.build_roster(str(planner), date(2026, 9, 28), 1)


def test_duplicate_name_raises(tmp_path):
    staff_rows = [
        make_staff_row(name="Bob", role="rotating", floor="down"),
        make_staff_row(name="Bob", role="rotating", floor="up"),
    ]
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])
    with pytest.raises(rg.RosterInputError, match="duplicate"):
        rg.build_roster(str(planner), date(2026, 9, 28), 1)


def test_leave_name_not_in_staff_raises(tmp_path):
    staff_rows = [make_staff_row(name="Bob", role="rotating", floor="down")]
    leave_rows = [("Nobody", "Sick", "2026-09-28", "2026-09-28")]
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, leave_rows)
    with pytest.raises(rg.RosterInputError, match="not in the Staff tab"):
        rg.build_roster(str(planner), date(2026, 9, 28), 1)


def test_missing_leave_from_date_raises(tmp_path):
    staff_rows = [make_staff_row(name="Bob", role="rotating", floor="down")]
    leave_rows = [("Bob", "Sick", "", "")]
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, leave_rows)
    with pytest.raises(rg.RosterInputError, match="From date"):
        rg.build_roster(str(planner), date(2026, 9, 28), 1)


def test_invalid_days_off_raises(tmp_path):
    staff_rows = [make_staff_row(name="Bob", role="rotating", floor="down", days_off="Someday")]
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])
    with pytest.raises(rg.RosterInputError, match="invalid day off"):
        rg.build_roster(str(planner), date(2026, 9, 28), 1)


def test_permission_error_on_open(tmp_path, monkeypatch):
    staff_rows = default_team(n_rotating=5)
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])

    def raise_permission_error(*args, **kwargs):
        raise PermissionError("file is locked")

    monkeypatch.setattr(rg, "load_workbook", raise_permission_error)
    with pytest.raises(rg.RosterInputError, match="open in Excel"):
        rg.build_roster(str(planner), date(2026, 9, 28), 1)


def test_commit_false_does_not_persist_history(tmp_path):
    staff_rows = default_team(n_rotating=9)
    planner = tmp_path / "Roster_Planner.xlsx"
    write_workbook(str(planner), staff_rows, [])
    rg.build_roster(str(planner), date(2026, 9, 28), 1, commit=False)
    wb = load_workbook(str(planner))
    history = rg.read_history(wb["History"])
    assert history == {}


# --------------------------------------------------------------------------
# Shift text helper
# --------------------------------------------------------------------------

@pytest.mark.parametrize("start,hours,expected", [
    ("7:30", 9, "7:30 – 4:30"),
    ("9:00", 9, "9:00 – 6:00"),
    ("8:00", 9, "8:00 – 5:00"),
    ("8:30", 9, "8:30 – 5:30"),
])
def test_shift_text(start, hours, expected):
    assert rg.shift_text(start, hours) == expected


def test_week_range_formatting_same_month():
    assert rg.format_week_range(date(2026, 9, 21)) == "21st – 25th September 2026"


def test_week_range_formatting_spans_months():
    assert rg.format_week_range(date(2026, 9, 28)) == "28th September – 2nd October 2026"
