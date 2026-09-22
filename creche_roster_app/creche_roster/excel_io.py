"""Reading the input workbook and writing the roster back into it.

Input tabs (you edit these): Staff, Leave, Overrides, Settings, History.
Output tabs (generated, replaced on every run): Week 1..Week N, Checks, Totals.
"""

from __future__ import annotations

import re
import shutil
from datetime import date, timedelta
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from .layout import COL_WIDTHS_CHARS, NCOLS, STYLES, ordinal_runs, week_grid
from .models import (
    NDAYS,
    ROLES,
    SLOT_LABELS,
    SLOTS,
    InputError,
    Inputs,
    Leave,
    Override,
    Roster,
    Settings,
    Shift,
    Staff,
)
from .parsing import fmt_day, is_blank, norm_hours, parse_date, parse_time
from .sample import sample_inputs

FONT = "Arial"
GENERATED = re.compile(r"^(Week \d+|Checks|Totals)$")
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ==========================================================================
# Reading
# ==========================================================================
def _norm(v) -> str:
    return "" if v is None else str(v).strip()


def _data_rows(ws):
    """Yield (excel row number, values) for rows below the header that are not empty."""
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if any(not is_blank(v) for v in row):
            yield r, row


def _slot_from_value(v, shifts: Dict[str, Shift]) -> str:
    if isinstance(v, str):
        key = v.strip().lower().replace(" ", "")
        if key in SLOTS:
            return key
    t = parse_time(v)
    for sl, sh in shifts.items():
        if sh.start == t:
            return sl
    times = ", ".join(f"{shifts[sl].start:%H:%M}" for sl in SLOTS)
    raise ValueError(f"{t:%H:%M} is not one of the four start times ({times})")


def read_inputs(path) -> Inputs:
    wb = load_workbook(path, data_only=True)
    sheets = {ws.title.strip().lower(): ws for ws in wb.worksheets}
    problems: List[str] = []
    for need in ("staff", "settings"):
        if need not in sheets:
            problems.append(f"The workbook has no '{need.title()}' tab.")
    if problems:
        raise InputError(problems)

    # ---- Settings ------------------------------------------------------
    raw: Dict[str, object] = {}
    for _, row in _data_rows(sheets["settings"]):
        if not is_blank(row[0]):
            raw[_norm(row[0]).lower()] = row[1] if len(row) > 1 else None

    def get(key, default=None, required=False):
        v = raw.get(key)
        if is_blank(v):
            if required:
                problems.append(f"Settings: '{key}' is missing.")
            return default
        return v

    title = _norm(get("title", get("site_name", "STAFF ROSTER")))
    break_text = _norm(get("break_text", "10 MINS"))
    day_headers = _norm(get("day_headers", "no")).lower() in ("yes", "y", "true", "1")
    roster_start = None
    v = get("roster_start", required=True)
    if v is not None:
        try:
            roster_start = parse_date(v)
        except ValueError as e:
            problems.append(f"Settings roster_start: {e}")
    weeks = 4
    v = get("weeks", 4)
    try:
        weeks = int(v)
    except (TypeError, ValueError):
        problems.append(f"Settings weeks: '{v}' is not a whole number.")
    floors = [f.strip() for f in _norm(get("floors", "All")).split(",") if f.strip()]

    shifts: Dict[str, Shift] = {}
    for sl in SLOTS:
        try:
            a = get(f"{sl}_start", required=True)
            b = get(f"{sl}_end", required=True)
            if a is not None and b is not None:
                shifts[sl] = Shift(parse_time(a), parse_time(b))
        except ValueError as e:
            problems.append(f"Settings {sl}: {e}")

    min_open: Dict[str, int] = {}
    min_close: Dict[str, int] = {}
    for f in floors:
        slug = f.lower().replace(" ", "_")
        for table, prefix in ((min_open, "min_open"), (min_close, "min_close")):
            v = get(f"{prefix}_{slug}", required=True)
            if v is not None:
                try:
                    table[f] = int(v)
                except (TypeError, ValueError):
                    problems.append(f"Settings {prefix}_{slug}: '{v}' is not a whole number.")
    fallback_closer = _norm(get("fallback_closer", "")) or None
    if problems or roster_start is None:
        raise InputError(problems or ["Settings: roster_start is missing."])
    settings = Settings(
        title, roster_start, weeks, floors, shifts, min_open, min_close,
        break_text, day_headers, fallback_closer,
    )

    # ---- Staff -----------------------------------------------------------
    staff: List[Staff] = []
    canon: Dict[str, str] = {}
    floor_lookup = {f.lower(): f for f in floors}
    for r, row in _data_rows(sheets["staff"]):
        padded = row + (None,) * 12
        name, floor, role, note = padded[:4]
        day_vals = padded[4:4 + NDAYS]
        number = _norm(padded[9])
        start_raw, end_raw = padded[10], padded[11]
        name, role = _norm(name), _norm(role).lower()
        floor_in = _norm(floor)
        if role not in ROLES:
            problems.append(f"Staff row {r}: role '{role}' must be one of: {', '.join(ROLES)}.")
            continue
        if not floor_in and role in ("static", "vacant", "blank"):
            floor_ok = floors[0]  # these rows do not count towards cover, so the floor is irrelevant
        else:
            floor_ok = floor_lookup.get(floor_in.lower())
            if floor_ok is None:
                problems.append(f"Staff row {r}: floor '{floor_in}' is not one of: {', '.join(floors)}.")
                continue
        fixed_slot = None
        if role == "fixed":
            try:
                fixed_slot = "early" if is_blank(note) else _slot_from_value(note, shifts)
            except ValueError as e:
                problems.append(f"Staff row {r} ({name}): {e}")
        note_text = "" if is_blank(note) else (f"{note:%H:%M}" if hasattr(note, "hour") else _norm(note))
        if role in ("static", "vacant"):
            note_text = norm_hours(note_text)
        hours = [norm_hours(f"{v:%H:%M}" if hasattr(v, "hour") else v) for v in day_vals]
        if number and number != "-" and not number.isdigit():
            problems.append(f"Staff row {r}: 'Print no.' must be a number, or - for none.")
            number = ""
        start_date = end_date = None
        if not is_blank(start_raw):
            try:
                start_date = parse_date(start_raw)
            except ValueError as e:
                problems.append(f"Staff row {r} ({name}): Start date - {e}")
        if not is_blank(end_raw):
            try:
                end_date = parse_date(end_raw)
            except ValueError as e:
                problems.append(f"Staff row {r} ({name}): End date - {e}")
        if start_date and end_date and end_date < start_date:
            problems.append(f"Staff row {r} ({name}): End date is before Start date.")
        staff.append(Staff(name, floor_ok, role, note_text, fixed_slot, hours, number, start_date, end_date))
        if role not in ("vacant", "blank") and name:
            canon[name.lower()] = name

    def canonical(name_raw, sheet, r):
        n = _norm(name_raw)
        c = canon.get(n.lower())
        if c is None:
            problems.append(f"{sheet} row {r}: '{n}' is not in the Staff tab (check the spelling).")
        return c

    # ---- Leave ------------------------------------------------------------
    leave: List[Leave] = []
    if "leave" in sheets:
        for r, row in _data_rows(sheets["leave"]):
            name, frm, to, kind = (row + (None,) * 4)[:4]
            c = canonical(name, "Leave", r)
            try:
                start = parse_date(frm)
                end = None if is_blank(to) else parse_date(to)
            except ValueError as e:
                problems.append(f"Leave row {r}: {e}")
                continue
            if c:
                leave.append(Leave(c, start, end, _norm(kind) or "Leave"))

    # ---- Overrides ----------------------------------------------------------
    overrides: List[Override] = []
    if "overrides" in sheets:
        for r, row in _data_rows(sheets["overrides"]):
            name, frm, to, st = (row + (None,) * 4)[:4]
            c = canonical(name, "Overrides", r)
            try:
                start = parse_date(frm)
                end = None if is_blank(to) else parse_date(to)
                role = next((x.role for x in staff if x.name == c), None)
                if role == "static":
                    if is_blank(st):
                        raise ValueError("type the hours (e.g. 10:00 - 2:00) or OFF")
                    text = norm_hours(f"{st:%H:%M}" if hasattr(st, "hour") else st)
                    ov = Override(c, start, end, None, text)
                else:
                    ov = Override(c, start, end, _slot_from_value(st, shifts))
            except ValueError as e:
                problems.append(f"Overrides row {r}: {e}")
                continue
            if c:
                overrides.append(ov)

    # ---- History (unknown names are ignored: people leave) ----------------------
    history: Dict[str, Dict[str, int]] = {}
    last_slot: Dict[str, str] = {}
    if "history" in sheets:
        for r, row in _data_rows(sheets["history"]):
            row = (row + (None,) * 6)[:6]
            c = canon.get(_norm(row[0]).lower())
            if c is None:
                continue
            counts = {}
            for sl, v in zip(SLOTS, row[1:5]):
                try:
                    counts[sl] = 0 if is_blank(v) else int(v)
                except (TypeError, ValueError):
                    problems.append(f"History row {r}: '{v}' is not a whole number.")
            history[c] = counts
            if not is_blank(row[5]):
                try:
                    last_slot[c] = _slot_from_value(row[5], shifts)
                except ValueError as e:
                    problems.append(f"History row {r}: {e}")

    if problems:
        raise InputError(problems)
    return Inputs(settings=settings, staff=staff, leave=leave, overrides=overrides, history=history, last_slot=last_slot)


# ==========================================================================
# Writing the roster
# ==========================================================================
def _font(spec) -> Font:
    name = "Cambria" if spec.get("serif") else FONT
    return Font(name=name, size=spec["size"], bold=spec["bold"], italic=spec["italic"], color=spec["color"])


def _apply(cell, style_name: str) -> None:
    spec = STYLES[style_name]
    cell.font = _font(spec)
    if spec["fill"]:
        cell.fill = PatternFill("solid", start_color=spec["fill"], end_color=spec["fill"])
    cell.alignment = Alignment(horizontal=spec["align"], vertical="center", wrap_text=True)
    if spec["border"]:
        cell.border = BOX


def _subtitle_rich(text: str, spec) -> CellRichText:
    """'21st - 25th September 2026' with the st/nd/rd/th raised, as on the old roster."""
    base = dict(rFont=FONT, sz=spec["size"], b=spec["bold"])
    parts = []
    for run, raised in ordinal_runs(text):
        font = InlineFont(vertAlign="superscript", **base) if raised else InlineFont(**base)
        parts.append(TextBlock(font, run))
    return CellRichText(*parts)


def _write_week(ws, roster: Roster, wi: int) -> None:
    rows = week_grid(roster, wi)
    ws.sheet_view.showGridLines = False
    for c, w in enumerate(COL_WIDTHS_CHARS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    for r, row in enumerate(rows, start=1):
        ws.row_dimensions[r].height = 20 * row.height
        if row.merged:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NCOLS)
            lc = row.cells[0]
            spec = STYLES[lc.style]
            cell = ws.cell(r, 1)
            if lc.style == "subtitle":
                cell.value = _subtitle_rich(lc.text, spec)
            else:
                cell.value = lc.text if lc.text != "" else None
            _apply(cell, lc.style)
            if spec["border"] or spec["fill"]:
                for c in range(2, NCOLS + 1):
                    _apply(ws.cell(r, c), lc.style)
        else:
            for c, lc in enumerate(row.cells, start=1):
                cell = ws.cell(r, c, lc.text if lc.text != "" else None)
                _apply(cell, lc.style)

    last = len(rows)
    ws.print_area = f"A1:{get_column_letter(NCOLS)}{last}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6


def _write_checks(ws, roster: Roster) -> None:
    ws.sheet_view.showGridLines = False
    widths = [12, 10, 16, 22, 100]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    n_breach = len(roster.breaches)
    ws["A1"] = "Checks"
    ws["A1"].font = Font(name=FONT, size=16, bold=True)
    ws["A2"] = (
        "Hard-rule breaches: 0. Every opening, closing and pairing rule is met."
        if n_breach == 0
        else f"Hard-rule breaches: {n_breach}. See the red rows below."
    )
    ws["A2"].font = Font(name=FONT, size=11, bold=True, color="006100" if n_breach == 0 else "9C0006")

    r = 4
    for i, h in enumerate(["Level", "Week", "Day", "Rule", "Message"], start=1):
        c = ws.cell(r, i, h)
        _apply(c, "header")
    order = {"BREACH": 0, "WARNING": 1, "INFO": 2}
    items = sorted(roster.checks, key=lambda c: (order[c.level], c.week or 0, c.day or date.min))
    for chk in items:
        r += 1
        vals = [chk.level, (chk.week + 1) if chk.week is not None else "", fmt_day(chk.day) if chk.day else "", chk.rule, chk.message]
        style = {"BREACH": "note_bad", "WARNING": "note_warn", "INFO": "note"}[chk.level]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(r, i, v)
            _apply(c, style)
            c.border = BOX
            c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="left")
    if not items:
        r += 1
        ws.cell(r, 1, "Nothing to report.").font = Font(name=FONT, size=10, italic=True)


def _write_totals(ws, roster: Roster) -> None:
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([24, 10, 10, 10, 10, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["A1"] = "Totals"
    ws["A1"].font = Font(name=FONT, size=16, bold=True)
    r = 2

    def table(title: str, note: str, data: Dict[str, Dict[str, int]], with_last: bool) -> None:
        nonlocal r
        r += 2
        ws.cell(r, 1, title).font = Font(name=FONT, size=12, bold=True)
        r += 1
        ws.cell(r, 1, note).font = Font(name=FONT, size=9, italic=True, color="404040")
        r += 1
        headers = ["Name"] + [SLOT_LABELS[sl] for sl in SLOTS] + (["Last slot"] if with_last else [])
        for i, h in enumerate(headers, start=1):
            _apply(ws.cell(r, i, h), "header")
        for _, st in roster.staff_keys:
            if st.name not in data:
                continue
            r += 1
            _apply(ws.cell(r, 1, st.name), "plain")
            for i, sl in enumerate(SLOTS, start=2):
                _apply(ws.cell(r, i, data[st.name].get(sl, 0)), "plain")
            if with_last:
                ls = roster.last_slot.get(st.name)
                _apply(ws.cell(r, 6, SLOT_LABELS[ls] if ls else ""), "plain")

    table(
        "Days on each start time in this roster",
        "Working days only. Leave is not counted.",
        roster.period_counts,
        False,
    )
    table(
        "Running totals: copy into the History tab before the next run",
        "History plus this roster. Pasting these into History keeps the rotation fair from one run to the next.",
        roster.cumulative_counts,
        True,
    )


def write_roster_only(roster: Roster, output_path) -> Path:
    """A clean, printable workbook with just the roster weeks - no input
    tabs (Staff/Leave/Overrides/Settings/History) and no Checks/Totals.
    For handing someone the roster itself (e.g. the app's Print button),
    as distinct from write_workbook's full editable planner file."""
    output_path = Path(output_path)
    wb = Workbook()
    wb.remove(wb.active)
    for wi in range(len(roster.weeks)):
        _write_week(wb.create_sheet(f"Week {wi + 1}"), roster, wi)
    wb.save(output_path)
    return output_path


def write_workbook(input_path, roster: Roster, output_path=None, backup: bool = True) -> Path:
    input_path = Path(input_path)
    output_path = Path(output_path) if output_path else input_path
    if backup and output_path == input_path and input_path.exists():
        shutil.copy2(input_path, str(input_path) + ".bak")
    wb = load_workbook(input_path)
    for ws in list(wb.worksheets):
        if GENERATED.match(ws.title):
            wb.remove(ws)
    for wi in range(len(roster.weeks)):
        _write_week(wb.create_sheet(f"Week {wi + 1}"), roster, wi)
    _write_checks(wb.create_sheet("Checks"), roster)
    _write_totals(wb.create_sheet("Totals"), roster)
    wb.save(output_path)
    return output_path


# ==========================================================================
# Starter workbook
# ==========================================================================
def _next_monday(today: date) -> date:
    return today + timedelta(days=(7 - today.weekday()) % 7)


def _style_input_sheet(ws, headers: List[str], widths: List[int]) -> None:
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(1, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="305496", end_color="305496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def _input_font(cell) -> None:
    cell.font = Font(name=FONT, size=10, color="0000FF")  # blue = you type here


def make_template(path, start: Optional[date] = None) -> Path:
    path = Path(path)
    start = start or _next_monday(date.today())
    inp = sample_inputs(start)
    s = inp.settings
    wb = Workbook()

    # ---- Read me ----------------------------------------------------------------
    ws = wb.active
    ws.title = "Read me"
    lines = [
        ("Creche roster generator", True),
        ("", False),
        ("Blue text = cells you type into. Everything on the Week, Checks and Totals tabs is generated: never edit those.", False),
        ("", False),
        ("Staff      Name, Floor, Role, Fixed slot / hours, then optional own hours for Mon..Fri (static staff). Type OFF for a day off. 'Print no.': empty = automatic, - = no number.", True),
        ("  Roles: rotating (shares the four start times), fixed (always the same slot, e.g. early), paired (the two staff whose", False),
        ("  starts complement each other), static (own hours typed in), vacant (an empty post: hours optional, name empty), blank (a spacer row).", False),
        ("Leave      Name, From, To, Type. Leave 'To' empty if you do not know the return date. Type shows on the roster: Holiday, Maternity Leave, OFF ...", True),
        ("Overrides  Name, From, To, Start time. Sets someone's start time by hand (static staff: type their hours or OFF). 'To' empty = open ended.", True),
        ("Settings   Title, first Monday, number of weeks, break text, floors, the four start/end times, minimum opening and closing cover.", True),
        ("History    Days each person has already worked on each start time, so rotation stays fair between runs.", True),
        ("", False),
        ("Pairing rule: the two 'paired' staff are always complementary. One on early (07:30) means the other is on late (09:00).", True),
        ("When one of them is away, the pairing is suspended. Set the other person's start time on the Overrides tab.", False),
        ("The override only applies on days the partner is away. When the partner is back, the pairing takes over automatically.", False),
        ("", False),
        ("PLACEHOLDERS: the floor grouping and the minimum opening / closing cover on the Settings tab are guesses. Set the real ones.", True),
        ("History is seeded from the printed roster of 21st-25th September 2026, so the rotation continues from it.", False),
        ("", False),
        ("To generate:  python -m creche_roster build Roster_Planner.xlsx   (writes Week 1..N, Checks and Totals, plus a PDF)", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(i, 1, text)
        c.font = Font(name=FONT, size=14 if i == 1 else 10, bold=bold)
    ws.column_dimensions["A"].width = 130

    # ---- Staff --------------------------------------------------------------------
    ws = wb.create_sheet("Staff")
    _style_input_sheet(
        ws,
        ["Name", "Floor", "Role", "Fixed slot / hours", "Mon", "Tue", "Wed", "Thu", "Fri", "Print no.", "Start date", "End date"],
        [20, 12, 12, 22, 14, 14, 14, 14, 14, 10, 14, 14],
    )
    for r, st in enumerate(inp.staff, start=2):
        vals = [st.name or None, None if (st.role in ("static", "vacant", "blank")) else st.floor, st.role, st.note or None]
        vals += [h or None for h in st.hours]
        vals.append(st.number or None)
        vals += [st.start_date, st.end_date]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            _input_font(cell)
            if c in (11, 12):
                cell.number_format = "DD/MM/YYYY"
    dv_floor = DataValidation(type="list", formula1='"' + ",".join(s.floors) + '"', allow_blank=True)
    dv_role = DataValidation(type="list", formula1='"' + ",".join(ROLES) + '"', allow_blank=True)
    ws.add_data_validation(dv_floor)
    ws.add_data_validation(dv_role)
    dv_floor.add("B2:B60")
    dv_role.add("C2:C60")

    # ---- Leave ----------------------------------------------------------------------
    ws = wb.create_sheet("Leave")
    _style_input_sheet(ws, ["Name", "From", "To (empty = until back)", "Type"], [24, 14, 24, 18])
    for r, lv in enumerate(inp.leave, start=2):
        for c, v in enumerate([lv.name, lv.start, lv.end, lv.kind], start=1):
            cell = ws.cell(r, c, v)
            _input_font(cell)
            if c in (2, 3):
                cell.number_format = "DD/MM/YYYY"

    # ---- Overrides -------------------------------------------------------------------
    ws = wb.create_sheet("Overrides")
    _style_input_sheet(ws, ["Name", "From", "To (empty = until back)", "Start time (or own hours / OFF)"], [24, 14, 24, 30])
    for r, ov in enumerate(inp.overrides, start=2):
        for c, v in enumerate([ov.name, ov.start, ov.end, s.shifts[ov.slot].start if ov.slot else ov.text], start=1):
            cell = ws.cell(r, c, v)
            _input_font(cell)
            if c in (2, 3):
                cell.number_format = "DD/MM/YYYY"
            if c == 4:
                cell.number_format = "HH:MM"

    # ---- Settings ------------------------------------------------------------------------
    ws = wb.create_sheet("Settings")
    _style_input_sheet(ws, ["Setting", "Value", "What it does"], [26, 22, 90])
    rows = [
        ("title", s.title, "Big red title at the top of every printed week."),
        ("roster_start", s.roster_start, "First Monday of the roster."),
        ("weeks", s.weeks, "How many weeks to generate (1 to 12)."),
        ("break_text", s.break_text, "Text in the break column next to each person."),
        ("day_headers", "no" if not s.day_headers else "yes", "yes = print Mon..Fri headings. The old roster has none."),
        ("floors", ", ".join(s.floors), "Comma separated. Each floor needs its own min_open_ / min_close_ rows below. 'All' = one group."),
        ("fallback_closer", s.fallback_closer or "", "A static-hours person (matching a Staff name) who closes on a day neither paired person does. Leave blank for none."),
    ]
    for sl in SLOTS:
        rows.append((f"{sl}_start", s.shifts[sl].start, f"{SLOT_LABELS[sl]} start time."))
        rows.append((f"{sl}_end", s.shifts[sl].end, f"{SLOT_LABELS[sl]} end time."))
    for f in s.floors:
        slug = f.lower().replace(" ", "_")
        rows.append((f"min_open_{slug}", s.min_open[f], f"PLACEHOLDER: minimum staff on the early start ({f})."))
        rows.append((f"min_close_{slug}", s.min_close[f], f"PLACEHOLDER: minimum staff on the late start ({f})."))
    for r, (k, v, note) in enumerate(rows, start=2):
        ws.cell(r, 1, k).font = Font(name=FONT, size=10)
        c = ws.cell(r, 2, v)
        _input_font(c)
        if isinstance(v, date):
            c.number_format = "DD/MM/YYYY"
        elif hasattr(v, "hour"):
            c.number_format = "HH:MM"
        ws.cell(r, 3, note).font = Font(name=FONT, size=9, italic=True, color="404040")

    # ---- History ----------------------------------------------------------------------------
    ws = wb.create_sheet("History")
    _style_input_sheet(ws, ["Name", "Early", "Mid 1", "Mid 2", "Late", "Last slot"], [24, 10, 10, 10, 10, 12])
    people = [x for x in inp.staff if x.role in ("rotating", "fixed", "paired")]
    for r, st in enumerate(people, start=2):
        _input_font(ws.cell(r, 1, st.name))
        counts = inp.history.get(st.name, {})
        for c, sl in enumerate(SLOTS, start=2):
            _input_font(ws.cell(r, c, counts.get(sl, 0)))
        ls = inp.last_slot.get(st.name)
        _input_font(ws.cell(r, 6, SLOT_LABELS[ls] if ls else None))

    wb.save(path)
    return path
